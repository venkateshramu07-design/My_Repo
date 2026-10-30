# Submit the blog's addresses to Bing with the Bing Webmaster URL Submission API (needs BING_API_KEY in the environment).
# Sends the addresses not yet submitted, in the priority order of URL_INSPECTION.xlsx and then the sitemap, up to the
# site's daily quota, and records what was sent in bing_submitted.json so the next run carries on.
#   python3 seo/bing_submit.py            submit up to today's quota
#   python3 seo/bing_submit.py URL ...    submit just these (for example right after publishing)
import json, os, re, sys, time, urllib.request
HERE = os.path.dirname(os.path.abspath(__file__)); SITE = "https://randomthoughtsever.blogspot.com"
KEY = os.environ.get("BING_API_KEY") or sys.exit("BING_API_KEY is not set")
API = "https://ssl.bing.com/webmaster/api.svc/json/"
LOG = os.path.join(HERE, "bing_submitted.json")

def call(method, body=None, **q):
    q["apikey"] = KEY; url = API + method + "?" + "&".join("%s=%s" % (k, urllib.parse.quote(v, safe="")) for k, v in q.items())
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, headers={"Content-Type": "application/json; charset=utf-8"})
    try:
        return json.loads(urllib.request.urlopen(req, timeout=60).read() or b"{}")
    except urllib.error.HTTPError as e:
        sys.exit("%s failed: %s %s" % (method, e.code, e.read()[:300]))

import urllib.parse
quota = call("GetUrlSubmissionQuota", siteUrl=SITE).get("d", {})
daily = int(quota.get("DailyQuota") or 0); print("Bing quota today:", daily, "| this month:", quota.get("MonthlyQuota"))
done = json.load(open(LOG)) if os.path.exists(LOG) else {}
if len(sys.argv) > 1:
    todo = sys.argv[1:]
else:
    order = []
    try:
        from openpyxl import load_workbook
        ws = load_workbook(os.path.join(HERE, "URL_INSPECTION.xlsx"))["URL Inspection"]
        order = [ws.cell(r, 4).value for r in range(2, ws.max_row + 1)]
    except Exception:
        pass
    sm = urllib.request.urlopen(urllib.request.Request(SITE + "/sitemap.xml", headers={"User-Agent": "Mozilla/5.0"})).read().decode()
    order += [u for u in re.findall(r"<loc>([^<]+)</loc>", sm) if u not in order]
    todo = [SITE + "/"] + [u for u in order if u and u not in done]
todo = [u for u in todo if u not in done or len(sys.argv) > 1][:daily]
for i in range(0, len(todo), 100):
    batch = todo[i:i + 100]
    call("SubmitUrlBatch", {"siteUrl": SITE, "urlList": batch})
    for u in batch: done[u] = time.strftime("%Y-%m-%d")
    json.dump(done, open(LOG, "w"), indent=1)
    print("submitted", len(batch))
print("sent today:", len(todo), "| total submitted so far:", len(done))
